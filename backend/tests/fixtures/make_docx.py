"""Regenerate sample.docx: python tests/fixtures/make_docx.py (from backend/)."""

from pathlib import Path

from docx import Document

doc = Document()
doc.add_heading("Quarterly Report", level=1)
p = doc.add_paragraph("Revenue grew ")
p.add_run("strongly").bold = True
p.add_run(" this ")
p.add_run("quarter").italic = True
p.add_run(".")
u = doc.add_paragraph()
u.add_run("Underlined note").underline = True
doc.add_heading("Highlights", level=2)
doc.add_paragraph("New customers", style="List Bullet")
doc.add_paragraph("Lower churn", style="List Bullet")
doc.add_paragraph("Plan budget", style="List Number")
doc.add_paragraph("Hire team", style="List Number")
doc.add_heading("Appendix", level=4)
table = doc.add_table(rows=1, cols=2)
table.rows[0].cells[0].text = "a"
doc.save(Path(__file__).with_name("sample.docx"))
